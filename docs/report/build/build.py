"""Rebuild the report body (Introduction → References) for MediSimplify, reusing the template's formatting.

Usage: python build.py template.docx figs_dir out.docx
Everything before the Introduction (cover, certification, title page) is left untouched. The output
screenshot slots (images) are kept as they are; only their captions are rewritten.
"""
import copy
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

from content import ABSTRACT, CHAPTERS, CONCLUSION, INTRO, REFERENCES, SCREENS

TEMPLATE, FIGS, OUT = sys.argv[1], Path(sys.argv[2]), sys.argv[3]
doc = Document(TEMPLATE)
body = doc.element.body
els = list(body.iterchildren())
assert "Introduction" in "".join(t.text or "" for t in els[53].iter(qn("w:t")))
assert els[-1].tag == qn("w:sectPr")
FINAL_SECT = els[-1]


# ------------------------------------------------------------------ low-level helpers
def ppr(i):
    p = copy.deepcopy(els[i].find(qn("w:pPr")))
    for s in p.findall(qn("w:sectPr")):
        p.remove(s)
    return p


def sectpr(i):
    return copy.deepcopy(els[i].find(qn("w:pPr")).find(qn("w:sectPr")))


def rpr(sz=None, b=False, i=False, font=None):
    r = OxmlElement("w:rPr")
    if font:
        f = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            f.set(qn(a), font)
        r.append(f)
    if b:
        r.append(OxmlElement("w:b"))
    if i:
        r.append(OxmlElement("w:i"))
    if sz:
        s = OxmlElement("w:sz"); s.set(qn("w:val"), str(sz)); r.append(s)
        s = OxmlElement("w:szCs"); s.set(qn("w:val"), str(sz)); r.append(s)
    return r


def run(text, **fmt):
    r = OxmlElement("w:r")
    r.append(rpr(**fmt))
    t = OxmlElement("w:t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    return r


TOKEN = re.compile(r"(\*\*.+?\*\*|//.+?//|`.+?`)")


def runs(text, sz=None, b=False, i=False):
    out = []
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**"):
            out.append(run(part[2:-2], sz=sz, b=True, i=i))
        elif part.startswith("//"):
            out.append(run(part[2:-2], sz=sz, b=b, i=True))
        elif part.startswith("`"):
            out.append(run(part[1:-1], sz=sz, b=b, i=i, font="Courier New"))
        else:
            out.append(run(part, sz=sz, b=b, i=i))
    return out


def para(p_pr, content=(), **fmt):
    p = OxmlElement("w:p")
    p.append(p_pr)
    for r in (runs(content, **fmt) if isinstance(content, str) else content):
        p.append(r)
    return p


def set_ppr(p_pr, tag, attrs, before=("w:rPr",)):
    """Insert or replace a pPr child, keeping schema order by placing it before the given elements."""
    old = p_pr.find(qn(tag))
    if old is not None:
        p_pr.remove(old)
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), v)
    for name in before:
        anchor = p_pr.find(qn(name))
        if anchor is not None:
            anchor.addprevious(e)
            return p_pr
    p_pr.append(e)
    return p_pr


BOOKMARK = [100]
TOC_ENTRIES = []


def bookmarked(p, level, text):
    BOOKMARK[0] += 1
    name = f"_TocMS{BOOKMARK[0]}"
    start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"), str(BOOKMARK[0])); start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), str(BOOKMARK[0]))
    p.find(qn("w:pPr")).addnext(start)
    p.append(end)
    TOC_ENTRIES.append((level, text, name))
    return p


def section_break(sect):
    p = OxmlElement("w:p")
    pp = OxmlElement("w:pPr")
    pp.append(sect)
    p.append(pp)
    return p


# ------------------------------------------------------------------ block builders
BODY_PPR = 73
CH_EXEMPLARS = [  # per chapter: h2, h3, bullet, letters, numbered
    {"h2": 72, "h3": None, "bullet": 112, "letters": 81, "numbered": 121},
    {"h2": 119, "h3": 148, "bullet": 135, "letters": 81, "numbered": 121},
    {"h2": 188, "h3": 194, "bullet": 246, "letters": 81, "numbered": 121},
    {"h2": 267, "h3": None, "bullet": 326, "letters": 81, "numbered": 121},
]
FIG_NO = {"n": 0}


def table(caption, header, rows, widths):
    out = [para(ppr(94), caption)]
    tbl = OxmlElement("w:tbl")
    tbl.append(copy.deepcopy(els[96].find(qn("w:tblPr"))))
    grid = OxmlElement("w:tblGrid")
    for w in widths:
        g = OxmlElement("w:gridCol"); g.set(qn("w:w"), str(w)); grid.append(g)
    tbl.append(grid)
    for r_i, row in enumerate([header] + rows):
        tr = OxmlElement("w:tr")
        trpr = OxmlElement("w:trPr")
        if r_i == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        cant = OxmlElement("w:cantSplit"); trpr.insert(0, cant)
        tr.append(trpr)
        for c_i, (text, w) in enumerate(zip(row, widths)):
            tc = OxmlElement("w:tc")
            tcpr = OxmlElement("w:tcPr")
            tw = OxmlElement("w:tcW"); tw.set(qn("w:w"), str(w)); tw.set(qn("w:type"), "dxa"); tcpr.append(tw)
            va = OxmlElement("w:vAlign"); va.set(qn("w:val"), "center"); tcpr.append(va)
            tc.append(tcpr)
            pp = OxmlElement("w:pPr")
            st = OxmlElement("w:pStyle"); st.set(qn("w:val"), "TableParagraph"); pp.append(st)
            sp = OxmlElement("w:spacing"); sp.set(qn("w:before"), "50"); sp.set(qn("w:after"), "50"); pp.append(sp)
            ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "100"); ind.set(qn("w:right"), "100"); pp.append(ind)
            if r_i == 0 or (len(widths) == 4 and c_i > 0):
                jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); pp.append(jc)
            tc.append(para(pp, text, sz=22, b=(r_i == 0)))
            tr.append(tc)
        tbl.append(tr)
    out.append(tbl)
    out.append(para(ppr(BODY_PPR)))
    return out


def figure(file, caption, width):
    holder = doc.add_paragraph()
    holder.add_run().add_picture(str(FIGS / file), width=Inches(width))
    p = holder._p
    p.getparent().remove(p)
    pp = OxmlElement("w:pPr")
    keep = OxmlElement("w:keepNext"); pp.append(keep)
    sp = OxmlElement("w:spacing"); sp.set(qn("w:before"), "200"); sp.set(qn("w:after"), "60"); pp.append(sp)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); pp.append(jc)
    old = p.find(qn("w:pPr"))
    if old is not None:
        p.remove(old)
    p.insert(0, pp)
    return [p, para(ppr(141), caption), para(ppr(BODY_PPR))]


def listing(caption, lines):
    out = []
    for k, line in enumerate(lines):
        pp = OxmlElement("w:pPr")
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F2F2F2"); pp.append(shd)
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:before"), "80" if k == 0 else "0"); sp.set(qn("w:after"), "80" if k == len(lines) - 1 else "0"); sp.set(qn("w:line"), "240"); sp.set(qn("w:lineRule"), "auto")
        pp.append(sp)
        ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "600"); ind.set(qn("w:right"), "300"); pp.append(ind)
        p = OxmlElement("w:p"); p.append(pp); p.append(run(line, sz=18, font="Courier New"))
        out.append(p)
    out.append(para(ppr(141), caption))
    out.append(para(ppr(BODY_PPR)))
    return out


def screens():
    out = [para(ppr(305), SCREENS["intro"])]

    def caption(i, pair):
        p = OxmlElement("w:p")
        p.append(ppr(i))
        p.append(run(pair[0], sz=20))
        r = OxmlElement("w:r"); r.append(rpr(sz=20)); r.append(OxmlElement("w:tab")); p.append(r)
        p.append(run(pair[1], sz=20))
        return p

    out += [els[306], caption(307, SCREENS["fig_a"]), els[308], caption(309, SCREENS["fig_b"]), para(ppr(310), SCREENS["fig1"]), para(ppr(311))]
    out += [els[312], caption(313, SCREENS["fig_c"]), els[314], caption(315, SCREENS["fig_d"]), para(ppr(316), SCREENS["fig2"])]
    return out


def chapter(index, title, blocks):
    ex = CH_EXEMPLARS[index]
    head_ppr = set_ppr(ppr(70), "w:outlineLvl", {"w:val": "0"})
    out = [bookmarked(para(head_ppr, title, b=True, sz=36), 1, title), para(ppr(71))]
    for block in blocks:
        kind = block[0]
        if kind == "h2":
            out.append(bookmarked(para(ppr(ex["h2"]), block[1]), 2, block[1]))
        elif kind == "h3":
            out.append(para(ppr(ex["h3"]), block[1]))
        elif kind == "h4":
            out.append(para(ppr(149), block[1]))
        elif kind == "p":
            out.append(para(ppr(BODY_PPR), block[1]))
        elif kind in ("bullets", "letters", "numbered"):
            src = ex["bullet" if kind == "bullets" else kind]
            out += [para(ppr(src), item, sz=24) for item in block[1]]
        elif kind == "table":
            out += table(*block[1:])
        elif kind == "figure":
            out += figure(*block[1:])
        elif kind == "listing":
            out += listing(*block[1:])
        elif kind == "eq":
            pp = ppr(141)
            set_ppr(pp, "w:spacing", {"w:before": "120", "w:after": "120"}, before=("w:ind", "w:jc", "w:rPr"))
            out.append(para(pp, block[1], i=True, sz=24))
        elif kind == "screens":
            out += screens()
        else:
            raise ValueError(kind)
    return out


def toc():
    sdt = OxmlElement("w:sdt")
    sdt.append(copy.deepcopy(els[67].find(qn("w:sdtPr"))))
    content = OxmlElement("w:sdtContent")

    def fld(kind):
        r = OxmlElement("w:r"); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), kind); r.append(f); return r

    def instr(text):
        r = OxmlElement("w:r"); t = OxmlElement("w:instrText"); t.text = text
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve"); r.append(t); return r

    for k, (level, text, name) in enumerate(TOC_ENTRIES):
        pp = OxmlElement("w:pPr")
        st = OxmlElement("w:pStyle"); st.set(qn("w:val"), "TOC2" if level == 1 else "TOC3"); pp.append(st)
        tabs = OxmlElement("w:tabs"); tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right"); tab.set(qn("w:leader"), "dot"); tab.set(qn("w:pos"), "9496"); tabs.append(tab); pp.append(tabs)
        sp = OxmlElement("w:spacing"); sp.set(qn("w:before"), "240" if level == 1 else "60"); pp.append(sp)
        ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "425" if level == 1 else "850"); pp.append(ind)
        p = OxmlElement("w:p"); p.append(pp)
        if k == 0:
            p.append(fld("begin")); p.append(instr(' TOC \\o "1-2" \\h \\z \\u ')); p.append(fld("separate"))
        h = OxmlElement("w:hyperlink"); h.set(qn("w:anchor"), name); h.set(qn("w:history"), "1")
        h.append(run(text, b=(level == 1)))
        r = OxmlElement("w:r"); r.append(OxmlElement("w:tab")); h.append(r)
        h.append(fld("begin")); h.append(instr(f" PAGEREF {name} \\h ")); h.append(fld("separate")); h.append(run(""))
        h.append(fld("end"))
        p.append(h)
        if k == len(TOC_ENTRIES) - 1:
            p.append(fld("end"))
        content.append(p)
    sdt.append(content)
    return sdt


# ------------------------------------------------------------------ assemble
new = []
new.append(para(ppr(53), "Introduction", b=True, sz=36))
new += [para(ppr(54), t) for t in INTRO]
new.append(els[58])  # section break: end of front matter + introduction (footers, numbering start)
new.append(para(ppr(59), "Abstract", b=True, sz=36))
new += [para(ppr(60), t) for t in ABSTRACT[:-1]]
new.append(para(ppr(64), ABSTRACT[-1]))
new.append(els[65])
new.append(para(ppr(66), "Contents", b=True, sz=36))
TOC_POS = len(new)
new.append(None)  # placeholder for the TOC, built after headings are known
new += [els[68], els[69]]

chapter_breaks = []
ch1 = sectpr(90)
default_footer = OxmlElement("w:footerReference")
default_footer.set(qn("w:type"), "default"); default_footer.set(qn("r:id"), "rId14")
ch1.insert(0, default_footer)  # page numbers on every chapter page, not only the first
for idx, (title, blocks) in enumerate(CHAPTERS):
    new += chapter(idx, title, blocks)
    new.append(section_break(ch1 if idx == 0 else sectpr(116)))

conc_ppr = ppr(332)
new.append(bookmarked(para(conc_ppr, "Conclusion"), 1, "Conclusion"))
new += [para(ppr(333), t) for t in CONCLUSION]
new.append(section_break(sectpr(337)))
new.append(bookmarked(para(ppr(338), "References"), 1, "References"))
for parts in REFERENCES:
    p = OxmlElement("w:p"); p.append(ppr(339))
    r = OxmlElement("w:r"); r.append(rpr(sz=24)); r.append(OxmlElement("w:tab")); p.append(r)
    for part in parts:
        if isinstance(part, tuple):
            p.append(run(part[1], sz=24, i=True))
        else:
            p.append(run(part, sz=24))
    new.append(p)
new[TOC_POS] = toc()

for el in els[53:-1]:
    body.remove(el)
for el in new:
    FINAL_SECT.addprevious(el)

# Ask Word to refresh the table of contents (page numbers) when the document is opened.
settings = doc.settings.element
upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true")
for name in ("w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars", "w:rsids"):
    anchor = settings.find(qn(name))
    if anchor is not None:
        anchor.addprevious(upd)
        break
else:
    settings.append(upd)

# Drop images that are no longer referenced (old diagrams); keep the screenshot slots.
xml = body.xml
used = set(re.findall(r'r:(?:embed|id)="(rId\d+)"', xml))
part = doc.part
for rid, rel in list(part.rels.items()):
    if rel.reltype.endswith("/image") and rid not in used:
        del part.rels[rid]

doc.save(OUT)
print("saved", OUT, "| toc entries:", len(TOC_ENTRIES))
